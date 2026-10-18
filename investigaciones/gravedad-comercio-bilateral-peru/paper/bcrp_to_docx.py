"""
Genera Paper-BCRP-Estudios-Economicos.docx segun las normas de autores de la
Revista Estudios Economicos del BCRP (verificadas en vivo, 2026-10-18):
espanol, A4, doble espacio, una cara, cuadros numerados consecutivamente
como "Cuadro N" (no "Tabla"), formato academico estandar (sin diseno de
marca: Times New Roman, sin colores de seccion).
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INPUT  = Path(__file__).parent / "Paper-BCRP-Estudios-Economicos.md"
OUTPUT = Path(__file__).parent / "Paper-BCRP-Estudios-Economicos.docx"

BLACK = RGBColor(0x00, 0x00, 0x00)
GREY  = RGBColor(0x40, 0x40, 0x40)
FUENTE = 'Times New Roman'


def set_cell_border_bottom(cell, size='8'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:color'), '000000')
    borders.append(bottom)
    tcPr.append(borders)


doc = Document()

section = doc.sections[0]
section.page_width  = Cm(21.0)   # A4
section.page_height = Cm(29.7)
section.left_margin   = Cm(3.0)
section.right_margin  = Cm(3.0)
section.top_margin    = Cm(3.0)
section.bottom_margin = Cm(3.0)

styles = doc.styles


def style_para(name, base, size, bold=False, color=BLACK,
               align=WD_ALIGN_PARAGRAPH.LEFT,
               space_before=0, space_after=0, italic=False,
               double_spaced=False):
    try:
        s = styles[name]
    except KeyError:
        s = styles.add_style(name, 1)
    s.base_style = styles[base]
    s.font.name = FUENTE
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = color
    s.font.italic = italic
    s.paragraph_format.alignment = align
    s.paragraph_format.space_before = Pt(space_before)
    s.paragraph_format.space_after  = Pt(space_after)
    if double_spaced:
        s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    return s


style_para('PaperTitle',  'Normal', 14, bold=True, color=BLACK,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
style_para('PaperAuthor', 'Normal', 11, bold=False, color=BLACK,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
style_para('PaperAffil',  'Normal', 10, bold=False, color=GREY, italic=True,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
style_para('SectionHdr',  'Normal', 12, bold=True, color=BLACK,
           align=WD_ALIGN_PARAGRAPH.LEFT, space_before=14, space_after=6)
style_para('SubSecHdr',   'Normal', 12, bold=True, color=BLACK, italic=True,
           align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4)
style_para('Body',        'Normal', 12, bold=False, color=BLACK,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=0, double_spaced=True)
style_para('Bullet',      'Normal', 12, bold=False, color=BLACK,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=0, double_spaced=True)
style_para('Caption',     'Normal', 11, bold=True, color=BLACK,
           align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4)


def bullet(text):
    p = doc.add_paragraph(style='Bullet')
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.first_line_indent = Pt(-14)
    p.add_run('-  ' + text)
    return p


def md_strip(t):
    t = re.sub(r'\*\*(.+?)\*\*', r'\1', t)
    t = re.sub(r'\*(.+?)\*', r'\1', t)
    t = re.sub(r'`(.+?)`', r'\1', t)
    t = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', t)  # enlaces -> solo el texto
    return t


with open(INPUT, encoding='utf-8') as fh:
    md = fh.read()

lines = md.split('\n')

# --- Portada: titulo, autor, afiliacion ---
title_line = md_strip(lines[0].lstrip('# ').strip())
author_idx = next(j for j, l in enumerate(lines[1:], start=1) if l.strip().startswith('**'))
author_line = md_strip(lines[author_idx].strip())
affil_idx = author_idx + 2  # linea en blanco, luego afiliacion
affil_line = md_strip(lines[affil_idx].strip()) if lines[affil_idx].strip().startswith('¹') else None

p = doc.add_paragraph(style='PaperTitle')
p.paragraph_format.space_before = Pt(36)
p.add_run(title_line)

p2 = doc.add_paragraph(style='PaperAuthor')
p2.add_run(author_line)

if affil_line:
    p3 = doc.add_paragraph(style='PaperAffil')
    p3.add_run(affil_line)
    i = affil_idx + 1
else:
    i = author_idx + 1

doc.add_paragraph()  # espacio antes del resumen

cuadro_re = re.compile(r'^\*\*(Cuadro \d+\..*)\*\*$')

while i < len(lines):
    line = lines[i].rstrip()
    s = line.strip()

    if re.match(r'^---+$', s):
        i += 1
        continue

    m_cuadro = cuadro_re.match(s)
    if m_cuadro:
        p = doc.add_paragraph(style='Caption')
        p.add_run(m_cuadro.group(1))
        i += 1
        continue

    if s.startswith('```'):
        j = i + 1
        code_lines = []
        while j < len(lines) and not lines[j].strip().startswith('```'):
            code_lines.append(lines[j])
            j += 1
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('\n'.join(code_lines))
        r.font.name = 'Cambria Math'
        r.font.size = Pt(11)
        r.font.italic = True
        i = j + 1
        continue

    if s.startswith('|'):
        rows = []
        j = i
        while j < len(lines) and lines[j].strip().startswith('|'):
            row = [c.strip() for c in lines[j].split('|')[1:-1]]
            if row and any(c for c in row):
                rows.append(row)
            j += 1
            if j < len(lines) and re.match(r'\|[\s\-:|]+\|', lines[j].strip()):
                j += 1
        if rows:
            col_count = len(rows[0])
            tbl = doc.add_table(rows=len(rows), cols=col_count)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            width = (Cm(21.0) - Cm(3.0 * 2)) / col_count
            for row_idx, row_data in enumerate(rows):
                for col_idx, text in enumerate(row_data):
                    cell = tbl.rows[row_idx].cells[col_idx]
                    cell.width = width
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after  = Pt(2)
                    run = p.add_run(md_strip(text))
                    run.font.name = FUENTE
                    run.font.size = Pt(10)
                    if row_idx == 0:
                        run.font.bold = True
                        set_cell_border_bottom(cell)
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
        i = j
        continue

    if not s:
        i += 1
        continue

    if line.startswith('## '):
        texto = md_strip(s[3:].strip())
        if texto.lower() in ('resumen', 'referencias'):
            p = doc.add_paragraph(style='SectionHdr')
            p.add_run(texto)
        else:
            p = doc.add_paragraph(style='SectionHdr')
            p.add_run(texto)
        i += 1
        continue

    if line.startswith('### '):
        p = doc.add_paragraph(style='SubSecHdr')
        p.add_run(md_strip(s[4:].strip()))
        i += 1
        continue

    if s.startswith('>'):
        p = doc.add_paragraph(style='Body')
        p.paragraph_format.left_indent = Pt(24)
        r = p.add_run(md_strip(s.lstrip('> ')))
        r.font.italic = True
        i += 1
        continue

    if re.match(r'^\s*[-*]\s+', s):
        t = re.sub(r'^\s*[-*]\s+', '', s)
        bullet(md_strip(t))
        i += 1
        continue

    if re.match(r'^\d+\.\s+', s):
        t = re.sub(r'^\d+\.\s+', '', s)
        p = doc.add_paragraph(style='Bullet')
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = Pt(-14)
        p.add_run(md_strip(t))
        i += 1
        continue

    p = doc.add_paragraph(style='Body')
    p.add_run(md_strip(s))
    i += 1

doc.save(str(OUTPUT))
print(f'DOCX: {OUTPUT}')
