"""
Genera Propuesta-Inversion-Publica-Exportaciones-Regional.docx desde el markdown.
Mismo patron que investigaciones/rca-shift-share-cptpp-china/propuesta/propuesta_to_docx.py
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INPUT  = Path(__file__).parent / "Propuesta-Inversion-Publica-Exportaciones-Regional.md"
OUTPUT = Path(__file__).parent / "Propuesta-Inversion-Publica-Exportaciones-Regional.docx"

NAVY  = RGBColor(0x1a, 0x30, 0x50)
GREY  = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


doc = Document()

section = doc.sections[0]
section.page_width  = Cm(21.0)
section.page_height = Cm(29.7)
section.left_margin   = Cm(2.8)
section.right_margin  = Cm(2.8)
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)

styles = doc.styles


def style_para(name, base, size, bold=False, color=BLACK,
               align=WD_ALIGN_PARAGRAPH.LEFT,
               space_before=0, space_after=6, italic=False):
    try:
        s = styles[name]
    except KeyError:
        s = styles.add_style(name, 1)
    s.base_style = styles[base]
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = color
    s.font.italic = italic
    s.paragraph_format.alignment = align
    s.paragraph_format.space_before = Pt(space_before)
    s.paragraph_format.space_after  = Pt(space_after)
    return s


style_para('PropTitle',    'Normal', 18, bold=True,  color=NAVY,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6)
style_para('PropSubtitle', 'Normal', 11, bold=False, color=GREY,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
style_para('PropMeta',     'Normal', 9,  bold=False, color=GREY,
           align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=3)
style_para('SectionHdr',   'Normal', 11, bold=True,  color=WHITE,
           align=WD_ALIGN_PARAGRAPH.LEFT,   space_before=12, space_after=4)
style_para('SubSecHdr',    'Normal', 10, bold=True,  color=NAVY,
           align=WD_ALIGN_PARAGRAPH.LEFT,   space_before=8,  space_after=3)
style_para('Body',         'Normal', 10, bold=False, color=BLACK,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=5)
style_para('Bullet',       'Normal', 10, bold=False, color=BLACK,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=3)


def bullet(text):
    p = doc.add_paragraph(style='Bullet')
    p.paragraph_format.left_indent = Pt(20)
    p.paragraph_format.first_line_indent = Pt(-14)
    p.add_run('•  ' + text)
    return p


def add_section_header(text):
    p = doc.add_paragraph(style='SectionHdr')
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.font.color.rgb = WHITE
    run.font.bold = True
    run.font.size = Pt(11)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '1A3050')
    pPr.append(shd)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after  = Pt(4)
    pPr2 = p2._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1A3050')
    pBdr.append(bottom)
    pPr2.append(pBdr)


def md_strip(t):
    t = re.sub(r'\*\*(.+?)\*\*', r'\1', t)
    t = re.sub(r'\*(.+?)\*', r'\1', t)
    t = re.sub(r'`(.+?)`', r'\1', t)
    return t


with open(INPUT, encoding='utf-8') as fh:
    md = fh.read()

lines = md.split('\n')

# Portada fija
p = doc.add_paragraph(style='PropTitle')
p.paragraph_format.space_before = Pt(36)
p.paragraph_format.space_after  = Pt(8)
r = p.add_run("INFRAESTRUCTURA DE TRANSPORTE Y COMPETITIVIDAD\nEXPORTADORA REGIONAL EN EL PERU")
r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = NAVY

p2 = doc.add_paragraph(style='PropSubtitle')
p2.paragraph_format.space_before = Pt(0)
p2.paragraph_format.space_after  = Pt(6)
r2 = p2.add_run("Propuesta de investigacion en diseno")
r2.font.size = Pt(12); r2.font.color.rgb = GREY

p3 = doc.add_paragraph(style='PropMeta')
p3.paragraph_format.space_before = Pt(0)
p3.paragraph_format.space_after  = Pt(3)
r3 = p3.add_run("Precedencia temporal entre inversion publica productiva (Invierte.pe) y "
                "actividad exportadora regional: un diseno de panel distrital, 2001-2026")
r3.font.size = Pt(9); r3.font.italic = True; r3.font.color.rgb = GREY

p4 = doc.add_paragraph(style='PropMeta')
p4.paragraph_format.space_before = Pt(8)
r4 = p4.add_run("Version 1.0  |  Octubre 2026  |  Proyecto Rastro")
r4.font.size = Pt(9)

p5 = doc.add_paragraph(style='PropMeta')
p5.paragraph_format.space_before = Pt(2)
r5 = p5.add_run("Investigador principal: Mg. Ricardo Antonio Cuba Alvan")
r5.font.size = Pt(9)

doc.add_page_break()

# Saltar portada interna del markdown hasta el primer --- tras los metadatos
i = 0
while i < len(lines):
    s = lines[i].strip()
    if s.startswith('|') or (s.startswith('**') and ('investigador' in s.lower() or 'estado' in s.lower() or 'versión' in s.lower() or 'línea' in s.lower())):
        i += 1
        continue
    if re.match(r'^---+$', s):
        i += 1
        break
    i += 1

while i < len(lines):
    line = lines[i].rstrip()
    s = line.strip()

    if re.match(r'^---+$', s):
        i += 1
        continue

    if s.startswith('```'):
        j = i + 1
        code_lines = []
        while j < len(lines) and not lines[j].strip().startswith('```'):
            code_lines.append(lines[j])
            j += 1
        p = doc.add_paragraph()
        r = p.add_run('\n'.join(code_lines))
        r.font.name = 'Consolas'
        r.font.size = Pt(9)
        i = j + 1
        continue

    if s.startswith('|'):
        rows = []
        j = i
        while j < len(lines) and lines[j].strip().startswith('|'):
            row = [c.strip() for c in lines[j].split('|')[1:-1]]
            if row and any(c for c in row):
                rows.append(row)
            j += 1
            if j < len(lines) and re.match(r'\|[\s\-:|]+\|', lines[j].strip()):
                j += 1
        if rows:
            col_count = len(rows[0])
            tbl = doc.add_table(rows=len(rows), cols=col_count)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            width = (Cm(21.0) - Cm(2.8 * 2)) / col_count
            for row_idx, row_data in enumerate(rows):
                for col_idx, text in enumerate(row_data):
                    cell = tbl.rows[row_idx].cells[col_idx]
                    cell.width = width
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after  = Pt(3)
                    run = p.add_run(md_strip(text))
                    run.font.size = Pt(9)
                    if row_idx == 0:
                        set_cell_bg(cell, '1A3050')
                        run.font.bold = True
                        run.font.color.rgb = WHITE
                    else:
                        set_cell_bg(cell, 'F4F6FA' if row_idx % 2 == 0 else 'FFFFFF')
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        i = j
        continue

    if not s:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        i += 1
        continue

    if line.startswith('# '):
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after  = Pt(6)
        r = p.add_run(md_strip(s[2:]))
        r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = NAVY
        i += 1
        continue

    if line.startswith('## '):
        add_section_header(md_strip(s[3:].strip()))
        i += 1
        continue

    if line.startswith('### '):
        p = doc.add_paragraph(style='SubSecHdr')
        r = p.add_run(md_strip(s[4:].strip()))
        r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = NAVY
        i += 1
        continue

    if s.startswith('>'):
        p = doc.add_paragraph(style='Body')
        p.paragraph_format.left_indent = Pt(20)
        r = p.add_run(md_strip(s.lstrip('> ')))
        r.font.italic = True
        i += 1
        continue

    if re.match(r'^\s*[-*]\s+', s):
        t = re.sub(r'^\s*[-*]\s+', '', s)
        bullet(md_strip(t))
        i += 1
        continue

    if re.match(r'^\d+\.\s+', s):
        t = re.sub(r'^\d+\.\s+', '', s)
        p = doc.add_paragraph(style='Bullet')
        p.paragraph_format.left_indent = Pt(20)
        p.paragraph_format.first_line_indent = Pt(-14)
        p.add_run(md_strip(t))
        i += 1
        continue

    p = doc.add_paragraph(style='Body')
    p.add_run(md_strip(s))
    i += 1

doc.save(str(OUTPUT))
print(f'DOCX: {OUTPUT}')
