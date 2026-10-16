"""
Genera Hinrich-Proposal-Governance-Trade-Peru-EN.docx desde el markdown
de la propuesta en formato Hinrich (5 secciones requeridas por los
lineamientos de la convocatoria). Parser general de markdown: títulos,
listas, tablas, citas y párrafos con negritas/cursivas.
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path(__file__).parent
INPUT = BASE / "Hinrich-Proposal-Governance-Trade-Peru-EN.md"
OUTPUT = BASE / "Hinrich-Proposal-Governance-Trade-Peru-EN.docx"

NAVY = RGBColor(0x1a, 0x30, 0x50)
GREY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_inline_runs(p, text):
    """Convierte **negrita** e *cursiva* en runs con formato."""
    for token in re.split(r'(\*\*.+?\*\*|\*.+?\*)', text):
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            r = p.add_run(token[2:-2]); r.font.bold = True
        elif token.startswith('*') and token.endswith('*'):
            r = p.add_run(token[1:-1]); r.font.italic = True
        else:
            p.add_run(token)


def add_section_header(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    r.font.color.rgb = WHITE
    r.font.bold = True
    r.font.size = Pt(11)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), '1A3050')
    pPr.append(shd)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(4)
    pPr2 = p2._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '1A3050')
    pBdr.append(bottom); pPr2.append(pBdr)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0); section.page_height = Cm(29.7)
section.left_margin = Cm(2.5); section.right_margin = Cm(2.5)
section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.2)

normal = doc.styles['Normal']
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)

lines = INPUT.read_text(encoding='utf-8').split('\n')
i = 0
first_h1 = True
while i < len(lines):
    line = lines[i].rstrip()
    s = line.strip()

    if not s:
        i += 1; continue

    if re.match(r'^---+$', s):
        i += 1; continue

    if s.startswith('|'):
        rows = []
        j = i
        while j < len(lines) and lines[j].strip().startswith('|'):
            row = [c.strip() for c in lines[j].split('|')[1:-1]]
            if row and any(c for c in row):
                rows.append(row)
            j += 1
            if j < len(lines) and re.match(r'\|[\s\-:|]+\|', lines[j].strip()):
                j += 1
        if rows:
            n_cols = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=n_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            width = (Cm(21.0) - Cm(5.0)) / n_cols
            for ri, row_data in enumerate(rows):
                for ci in range(n_cols):
                    cell = tbl.rows[ri].cells[ci]
                    cell.width = width
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    text = row_data[ci] if ci < len(row_data) else ''
                    if ri == 0:
                        set_cell_bg(cell, '1A3050')
                        r = p.add_run(re.sub(r'\*\*(.+?)\*\*', r'\1', text))
                        r.font.bold = True; r.font.color.rgb = WHITE; r.font.size = Pt(9)
                    else:
                        if ri % 2 == 0:
                            set_cell_bg(cell, 'F4F6FA')
                        add_inline_runs(p, text)
                        for r in p.runs:
                            r.font.size = Pt(9)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        i = j; continue

    if line.startswith('# '):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(re.sub(r'\*\*(.+?)\*\*', r'\1', s[2:]))
        r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = NAVY
        first_h1 = False
        i += 1; continue

    if line.startswith('## '):
        add_section_header(doc, re.sub(r'\*\*(.+?)\*\*', r'\1', s[3:]))
        i += 1; continue

    if line.startswith('### '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(re.sub(r'\*\*(.+?)\*\*', r'\1', s[4:]))
        r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = NAVY
        i += 1; continue

    if s.startswith('>'):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(20)
        r = p.add_run(re.sub(r'\*\*(.+?)\*\*', r'\1', s.lstrip('> ')))
        r.font.italic = True
        i += 1; continue

    if re.match(r'^\s*[-*]\s+', s):
        t = re.sub(r'^\s*[-*]\s+', '', s)
        p = doc.add_paragraph(style='List Bullet')
        add_inline_runs(p, t)
        i += 1; continue

    if re.match(r'^\d+\.\s+', s):
        t = re.sub(r'^\d+\.\s+', '', s)
        p = doc.add_paragraph(style='List Number')
        add_inline_runs(p, t)
        i += 1; continue

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline_runs(p, s)
    i += 1

doc.save(str(OUTPUT))
print(f'DOCX: {OUTPUT}')
