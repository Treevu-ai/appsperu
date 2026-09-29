import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\acuba\appsperu\data\sunat\cdro_15.xlsx', data_only=True)
ws = wb.active
print(f'Dims: {ws.max_row}r x {ws.max_column}c')

# Print rows 1-12
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i >= 12:
        break
    vals = [str(c)[:20] if c is not None else 'None' for c in row[:5]]
    print(f'  [{i}] ' + ' | '.join(vals))
