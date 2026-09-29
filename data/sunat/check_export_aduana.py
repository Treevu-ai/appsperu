import openpyxl, os

base = r'C:\Users\acuba\appsperu\data\sunat'

for fname in ['cdro_21.xlsx', 'cdro_22.xlsx']:
    path = os.path.join(base, fname)
    if not os.path.exists(path):
        continue
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    print(f'=== {fname} ({ws.max_row}r x {ws.max_column}c) ===')
    # Show first 12 rows headers
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= 12:
            break
        vals = [str(c)[:30] if c is not None else '' for c in row[:8]]
        print(f'  r{i+1}: ' + ' | '.join(vals))

    # Check for SALAVERRY anywhere
    found = False
    for ridx, row in enumerate(ws.iter_rows(values_only=True)):
        for cell in row:
            if cell and isinstance(cell, str) and 'salaverry' in cell.lower():
                ctx = [str(c)[:25] if c else '' for c in row[:7]]
                print('  SALAVERRY at r' + str(ridx+1) + ': ' + ' | '.join(ctx))
                found = True
                break
    if not found:
        print('  (no SALAVERRY found)')
    print()
